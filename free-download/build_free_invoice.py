#!/usr/bin/env python3
"""
Build the FREE "Simple Invoice Template" lead-magnet workbook.

Follows the exact precedent of build/free_tracker.py: a deliberately small,
free, no-signup subset that funnels interested users to the paid Freelancer
Client & Invoice Tracker (part of the Freelancer Finance Toolkit) and to the
custom-spreadsheet service.

3 sheets: "Start Here" (setup + Settings + "Want more?"), "Invoice" (a
printable single-page invoice), "Invoice Log" (a simple 30-row manual log
with small SUMIF totals).

Run with:    .venv/bin/python build/free_invoice.py
Verify with: .venv/bin/python build/free_invoice.py --verify

Follows SPEC.md: design system, Google Sheets compatibility rules, and the
shared Start Here sheet convention (via build/common.py).
"""
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT_PATH = ROOT / "dist" / "free" / "GrowthChief-Simple-Invoice-Template.xlsx"

PRODUCT_NAME = "Simple Invoice Template (Free)"

ETSY_TOOLKIT_URL = "https://growthchief.etsy.com/listing/4560433852"
ETSY_CUSTOM_URL = "https://growthchief.etsy.com/listing/4564738495"

link_font = Font(name=common.FONT_NAME, size=11, bold=True, color=common.ACCENT_COLOR, underline="single")
section_font = Font(name=common.FONT_NAME, size=13, bold=True, color=common.HEADER_FILL_COLOR)
label_bold_font = Font(name=common.FONT_NAME, size=11, bold=True, color="111827")

# ---------------------------------------------------------------------------
# Invoice sheet layout constants
# ---------------------------------------------------------------------------
ITEM_HEADER_ROW = 13
ITEM_FIRST_ROW = 14
ITEM_SAMPLE_ROWS = 4
ITEM_TOTAL_ROWS = 12
ITEM_LAST_ROW = ITEM_FIRST_ROW + ITEM_TOTAL_ROWS - 1  # 25
SUBTOTAL_ROW = 27
TAX_PCT_ROW = 28
TAX_AMT_ROW = 29
TOTAL_ROW = 30
INVOICE_NUM_ROW = 10
INVOICE_DATE_ROW = 10
DUE_DATE_ROW = 11

# Sample invoice line items
SAMPLE_ITEMS = [
    ("Website design -- homepage & 2 inner pages (example)", 1, 800.0),
    ("Logo design (example)", 1, 250.0),
    ("Stock photography licensing (example)", 3, 15.0),
    ("Revision hours (example)", 2, 60.0),
]

# ---------------------------------------------------------------------------
# Invoice Log sheet layout constants
# ---------------------------------------------------------------------------
LOG_HEADER_ROW = 4
LOG_FIRST_ROW = 5
LOG_TOTAL_ROWS = 30
LOG_LAST_ROW = LOG_FIRST_ROW + LOG_TOTAL_ROWS - 1  # 34
LOG_SAMPLE_ROWS = 3

STATUS_OPTIONS = ["Draft", "Sent", "Paid", "Overdue"]

SAMPLE_LOG_ROWS = [
    ("INV-1001", "Bright Studio (example -- replace me)", date(2026, 8, 20), date(2026, 9, 3), 1312.20, "Paid"),
    ("INV-1002", "Acme Co (example -- replace me)", date(2026, 8, 25), date(2026, 9, 8), 500.0, "Sent"),
    ("INV-1003", "Rivertown Cafe (example -- replace me)", date(2026, 8, 29), date(2026, 9, 12), 250.0, "Overdue"),
]


# ---------------------------------------------------------------------------
# Start Here
# ---------------------------------------------------------------------------

def build_start_here_sheet(wb):
    steps = [
        "Set your currency symbol, business name, and default payment terms in Settings below -- the Invoice tab picks them up automatically.",
        "Go to the Invoice tab and fill in the highlighted cells: your business details, who you're billing, the invoice number and date, and up to 12 line items.",
        "The Due Date and every total calculate themselves -- Qty x Unit Price per line, a Subtotal, a Tax amount from the Tax % you enter, and a Total.",
        "Print the Invoice tab or save it as a PDF to send to your client (File > Print > Save as PDF in Excel or Google Sheets).",
        "Log each invoice you send on the Invoice Log tab so you can see what's been paid and what's still outstanding at a glance.",
        "Works in Excel and Google Sheets as-is -- no macros, no add-ons, nothing to enable.",
    ]
    settings = [
        {
            "label": "Currency symbol",
            "value": "$",
            "name": "CurrencySymbol",
            "note": "Used in headers throughout this workbook. Doesn't convert currencies -- just changes the symbol shown.",
        },
        {
            "label": "Your business name",
            "value": "Your Business Name",
            "name": "YourBusinessName",
            "note": "Fills in the \"From\" block on the Invoice tab automatically. Change it once here.",
        },
        {
            "label": "Payment terms (days)",
            "value": 14,
            "name": "PaymentTermsDays",
            "note": "Default number of days clients get to pay. The Invoice tab's Due Date is Invoice Date plus this number.",
        },
    ]
    result = common.build_start_here(
        wb,
        PRODUCT_NAME,
        steps,
        settings,
        guide_link_placeholder=(
            "This is the free version -- it's simple by design, so there's no separate guide. "
            "If anything's unclear, the download page has a full walkthrough."
        ),
    )
    ws = result["ws"]
    row = result["settings_range"][1] + 2

    header = ws.cell(row=row, column=2, value="Want more?")
    header.font = section_font
    row += 1

    intro = ws.cell(
        row=row, column=2,
        value=(
            "This free template covers the basics: one printable invoice, plus a simple log to track "
            "what you've sent. The full Freelancer Finance Toolkit builds on top of it with:"
        ),
    )
    intro.font = common.body_font
    intro.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 32
    row += 1

    bullets = [
        "A full Client & Invoice Tracker: auto-numbered invoices, a Due Date that calculates itself, and a Status dropdown (Draft, Sent, Paid, Overdue, Cancelled).",
        "A live Dashboard showing your outstanding balance, overdue total and count, and your average days to payment.",
        "A Clients tab that totals what each client has been invoiced and paid automatically, plus a Pipeline tab for tracking leads and proposals with a weighted deal value.",
        "The rest of the Freelancer Finance Toolkit alongside it: a 12-month income & expense tracker with a category dashboard, a tax set-aside calculator, and a rate calculator.",
    ]
    for b in bullets:
        c = ws.cell(row=row, column=2, value=f"-  {b}")
        c.font = common.body_font
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 28
        row += 1

    row += 1
    link1 = ws.cell(row=row, column=2, value="See the full Freelancer Finance Toolkit on Etsy")
    link1.font = link_font
    link1.hyperlink = ETSY_TOOLKIT_URL
    row += 1
    link2 = ws.cell(row=row, column=2, value="Need something custom built for your business? Custom spreadsheets, made to order")
    link2.font = link_font
    link2.hyperlink = ETSY_CUSTOM_URL
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    row += 2

    return ws


# ---------------------------------------------------------------------------
# Invoice (printable single-page invoice)
# ---------------------------------------------------------------------------

def build_invoice_sheet(wb):
    ws = common.new_data_sheet(wb, "Invoice", tab_color=common.TAB_TEAL, gridlines=False)
    common.set_col_widths(ws, {"A": 30, "B": 12, "C": 16, "D": 18})

    common.style_title(ws, "A1", "INVOICE")
    sub = ws.cell(row=2, column=1, value=(
        "Fill in the highlighted cells, then print this tab or save it as a PDF "
        "(File > Print > Save as PDF) to send to your client."
    ))
    sub.font = common.subtitle_font
    sub.alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
    ws.row_dimensions[2].height = 18

    # --- From / Bill To ---
    ws.cell(row=4, column=1, value="From").font = section_font
    ws.cell(row=4, column=3, value="Bill To").font = section_font

    from_cell = ws.cell(row=5, column=1, value="=YourBusinessName")
    from_cell.font = label_bold_font
    from_addr = ws.cell(row=6, column=1, value="Your business address")
    from_email = ws.cell(row=7, column=1, value="you@example.com")
    from_phone = ws.cell(row=8, column=1, value="(555) 555-1234")
    for cell in (from_addr, from_email, from_phone):
        cell.fill = common.accent_light_fill
        cell.font = common.body_font

    bill_cell = ws.cell(row=5, column=3, value="Bright Studio (example -- replace me)")
    bill_cell.font = label_bold_font
    bill_cell.fill = common.accent_light_fill
    bill_addr = ws.cell(row=6, column=3, value="456 Sample Ave, Rivertown")
    bill_email = ws.cell(row=7, column=3, value="billing@brightstudio-example.com")
    bill_phone = ws.cell(row=8, column=3, value="(555) 555-6789")
    for cell in (bill_addr, bill_email, bill_phone):
        cell.fill = common.accent_light_fill
        cell.font = common.body_font

    for r in range(5, 9):
        for c in (1, 3):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)

    # --- Invoice #, Invoice Date, Due Date, Payment Terms ---
    ws.cell(row=INVOICE_NUM_ROW, column=1, value="Invoice #").font = label_bold_font
    inv_num = ws.cell(row=INVOICE_NUM_ROW, column=2, value="INV-1001")
    inv_num.fill = common.accent_light_fill
    inv_num.border = common.thin_border
    inv_num.alignment = Alignment(horizontal="center")

    ws.cell(row=INVOICE_DATE_ROW, column=3, value="Invoice Date").font = label_bold_font
    inv_date = ws.cell(row=INVOICE_DATE_ROW, column=4, value=date(2026, 9, 6))
    inv_date.number_format = common.DATE_FMT
    inv_date.fill = common.accent_light_fill
    inv_date.border = common.thin_border
    inv_date.alignment = Alignment(horizontal="center")

    ws.cell(row=DUE_DATE_ROW, column=1, value="Due Date").font = label_bold_font
    due_date = ws.cell(
        row=DUE_DATE_ROW, column=2,
        value=f"=IF(D{INVOICE_DATE_ROW}=\"\",\"\",D{INVOICE_DATE_ROW}+PaymentTermsDays)",
    )
    due_date.number_format = common.DATE_FMT
    due_date.border = common.thin_border
    due_date.alignment = Alignment(horizontal="center")

    ws.cell(row=DUE_DATE_ROW, column=3, value="Payment Terms (days)").font = label_bold_font
    terms = ws.cell(row=DUE_DATE_ROW, column=4, value="=PaymentTermsDays")
    terms.border = common.thin_border
    terms.alignment = Alignment(horizontal="center")

    # --- Line item table ---
    common.style_header_row(ws, ITEM_HEADER_ROW, 1, 4)
    headers = ["Description", "Qty", '="Unit Price (" & CurrencySymbol & ")"', '="Amount (" & CurrencySymbol & ")"']
    for col, val in enumerate(headers, start=1):
        ws.cell(row=ITEM_HEADER_ROW, column=col, value=val)

    for i in range(ITEM_TOTAL_ROWS):
        r = ITEM_FIRST_ROW + i
        if i < ITEM_SAMPLE_ROWS:
            desc, qty, price = SAMPLE_ITEMS[i]
            ws.cell(row=r, column=1, value=desc)
            ws.cell(row=r, column=2, value=qty)
            ws.cell(row=r, column=3, value=price)
        amt = ws.cell(
            row=r, column=4,
            value=f'=IF(B{r}="","",IF(C{r}="","",B{r}*C{r}))',
        )
        amt.number_format = common.NUMBER_FMT
        ws.cell(row=r, column=3).number_format = common.NUMBER_FMT

    common.band_rows(ws, ITEM_FIRST_ROW, ITEM_LAST_ROW, 1, 4)

    # --- Subtotal / Tax / Total ---
    amount_range = f"D{ITEM_FIRST_ROW}:D{ITEM_LAST_ROW}"

    ws.merge_cells(start_row=SUBTOTAL_ROW, start_column=1, end_row=SUBTOTAL_ROW, end_column=3)
    ws.cell(row=SUBTOTAL_ROW, column=1, value='="Subtotal (" & CurrencySymbol & ")"').font = label_bold_font
    subtotal = ws.cell(row=SUBTOTAL_ROW, column=4, value=f"=SUM({amount_range})")
    subtotal.number_format = common.NUMBER_FMT
    subtotal.border = common.thin_border

    ws.merge_cells(start_row=TAX_PCT_ROW, start_column=1, end_row=TAX_PCT_ROW, end_column=3)
    ws.cell(row=TAX_PCT_ROW, column=1, value="Tax %").font = label_bold_font
    tax_pct = ws.cell(row=TAX_PCT_ROW, column=4, value=0.08)
    tax_pct.number_format = common.PERCENT_FMT
    tax_pct.fill = common.accent_light_fill
    tax_pct.border = common.thin_border

    ws.merge_cells(start_row=TAX_AMT_ROW, start_column=1, end_row=TAX_AMT_ROW, end_column=3)
    ws.cell(row=TAX_AMT_ROW, column=1, value='="Tax Amount (" & CurrencySymbol & ")"').font = label_bold_font
    tax_amt = ws.cell(row=TAX_AMT_ROW, column=4, value=f"=ROUND(D{SUBTOTAL_ROW}*D{TAX_PCT_ROW},2)")
    tax_amt.number_format = common.NUMBER_FMT
    tax_amt.border = common.thin_border

    ws.merge_cells(start_row=TOTAL_ROW, start_column=1, end_row=TOTAL_ROW, end_column=3)
    total_label = ws.cell(row=TOTAL_ROW, column=1, value='="Total (" & CurrencySymbol & ")"')
    total_label.font = Font(name=common.FONT_NAME, size=13, bold=True, color=common.HEADER_FILL_COLOR)
    total_val = ws.cell(row=TOTAL_ROW, column=4, value=f"=ROUND(D{SUBTOTAL_ROW}+D{TAX_AMT_ROW},2)")
    total_val.number_format = common.NUMBER_FMT
    total_val.font = Font(name=common.FONT_NAME, size=14, bold=True, color=common.HEADER_FILL_COLOR)
    total_val.fill = common.accent_light_fill
    total_val.border = common.thin_border
    ws.row_dimensions[TOTAL_ROW].height = 26

    # --- Print area: single page, portrait ---
    ws.print_area = "A1:D31"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.sheet_view.zoomScale = 100
    return ws


# ---------------------------------------------------------------------------
# Invoice Log
# ---------------------------------------------------------------------------

def build_invoice_log_sheet(wb):
    ws = common.new_data_sheet(wb, "Invoice Log", tab_color=common.TAB_SLATE)
    common.set_col_widths(ws, {"A": 14, "B": 26, "C": 12, "D": 12, "E": 14, "F": 12})

    common.style_title(ws, "A1", "Invoice Log")
    sub = ws.cell(row=2, column=1, value=(
        "Log every invoice you send here -- Invoice #, Client, Date, Due date, Amount, and Status are "
        "all entered by hand. The totals below add themselves up."
    ))
    sub.font = common.subtitle_font
    sub.alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=6)
    ws.row_dimensions[2].height = 18

    common.style_header_row(ws, LOG_HEADER_ROW, 1, 6)
    headers = ["Invoice #", "Client", "Date", "Due", '="Amount (" & CurrencySymbol & ")"', "Status"]
    for col, val in enumerate(headers, start=1):
        ws.cell(row=LOG_HEADER_ROW, column=col, value=val)

    for i, (num, client, d, due, amt, status) in enumerate(SAMPLE_LOG_ROWS):
        r = LOG_FIRST_ROW + i
        ws.cell(row=r, column=1, value=num)
        ws.cell(row=r, column=2, value=client)
        ws.cell(row=r, column=3, value=d)
        ws.cell(row=r, column=4, value=due)
        ws.cell(row=r, column=5, value=amt)
        ws.cell(row=r, column=6, value=status)

    common.band_rows(ws, LOG_FIRST_ROW, LOG_LAST_ROW, 1, 6)

    for r in range(LOG_FIRST_ROW, LOG_LAST_ROW + 1):
        ws.cell(row=r, column=3).number_format = common.DATE_FMT
        ws.cell(row=r, column=4).number_format = common.DATE_FMT
        ws.cell(row=r, column=5).number_format = common.NUMBER_FMT

    status_range = f"F{LOG_FIRST_ROW}:F{LOG_LAST_ROW}"
    common.add_list_validation(ws, status_range, STATUS_OPTIONS)
    common.freeze_header(ws, f"A{LOG_FIRST_ROW}")

    # --- Small totals ---
    amount_range = f"E{LOG_FIRST_ROW}:E{LOG_LAST_ROW}"
    status_range_abs = f"F{LOG_FIRST_ROW}:F{LOG_LAST_ROW}"

    totals_row = LOG_LAST_ROW + 2
    ws.cell(row=totals_row, column=1, value="Totals").font = section_font
    totals_row += 1

    def _total(row, label, formula):
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = label_bold_font
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        vc = ws.cell(row=row, column=5, value=formula)
        vc.font = Font(name=common.FONT_NAME, size=12, bold=True, color=common.HEADER_FILL_COLOR)
        vc.fill = common.accent_light_fill
        vc.border = common.thin_border
        vc.number_format = common.NUMBER_FMT
        return vc

    _total(totals_row, '="Total Invoiced (" & CurrencySymbol & ")"', f"=SUM({amount_range})")
    totals_row += 1
    _total(totals_row, '="Total Paid (" & CurrencySymbol & ")"', f'=SUMIF({status_range_abs},"Paid",{amount_range})')
    totals_row += 1
    _total(totals_row, '="Total Outstanding (" & CurrencySymbol & ")"', f'=SUMIF({status_range_abs},"<>Paid",{amount_range})')

    ws.sheet_view.zoomScale = 100
    return ws


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------

def build():
    wb = Workbook()
    build_start_here_sheet(wb)
    build_invoice_sheet(wb)
    build_invoice_log_sheet(wb)

    assert wb.sheetnames == ["Start Here", "Invoice", "Invoice Log"], wb.sheetnames

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} ({OUT_PATH.stat().st_size:,} bytes)")
    return OUT_PATH


# ---------------------------------------------------------------------------
# Verify
# ---------------------------------------------------------------------------

TASK_ALLOWED_FUNCTIONS = {
    "SUM", "SUMIF", "SUMIFS", "COUNTIF", "IF", "IFERROR", "TODAY", "DATE",
    "TEXT", "EOMONTH", "ROUND",
}

_RANGE_RE = re.compile(r"(?<![!\w$'])(\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+)")


def _strip_strings(formula):
    return re.sub(r'"[^"]*"', "", formula)


def verify():
    assert OUT_PATH.exists(), f"{OUT_PATH} does not exist -- run without --verify first"
    wb = load_workbook(OUT_PATH)

    # 1. Sheet names, exactly 3, in order.
    assert wb.sheetnames == ["Start Here", "Invoice", "Invoice Log"], (
        f"unexpected sheet list: {wb.sheetnames}"
    )
    print("PASS: sheet names ==", wb.sheetnames)

    invoice = wb["Invoice"]
    log = wb["Invoice Log"]

    # 2. Formulas present on Invoice and Invoice Log.
    invoice_formula_cells = [
        c for row in invoice.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert len(invoice_formula_cells) >= 15, (
        f"expected >=15 formulas on Invoice, found {len(invoice_formula_cells)}"
    )
    print(f"PASS: Invoice has {len(invoice_formula_cells)} formula cell(s)")

    log_formula_cells = [
        c for row in log.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert len(log_formula_cells) >= 6, (
        f"expected >=6 formulas on Invoice Log, found {len(log_formula_cells)}"
    )
    print(f"PASS: Invoice Log has {len(log_formula_cells)} formula cell(s)")

    # Required specific formula cells.
    required_invoice_cells = [
        f"B{DUE_DATE_ROW}",  # Due Date
        f"D{SUBTOTAL_ROW}",  # Subtotal
        f"D{TAX_AMT_ROW}",   # Tax amount
        f"D{TOTAL_ROW}",     # Total
    ] + [f"D{r}" for r in range(ITEM_FIRST_ROW, ITEM_LAST_ROW + 1)]  # Amount = Qty*Price
    for ref in required_invoice_cells:
        v = invoice[ref].value
        assert isinstance(v, str) and v.startswith("="), f"Invoice!{ref} missing formula, got {v!r}"
    print("PASS: Due Date, all 12 line-item Amounts, Subtotal, Tax amount, and Total formulas present")

    # 3. Data validations: Status dropdown on Invoice Log.
    dvs = log.data_validations.dataValidation
    assert len(dvs) == 1, f"expected 1 data validation on Invoice Log, found {len(dvs)}"
    dv = dvs[0]
    assert dv.type == "list", f"unexpected validation type {dv.type!r}"
    assert f"F{LOG_FIRST_ROW}:F{LOG_LAST_ROW}" in str(dv.sqref), f"Status dropdown range missing: {dv.sqref}"
    print(f"PASS: 1 list data validation present on Invoice Log (Status): {dv.sqref}")

    # 4. No unprefixed cross-sheet ranges: every multi-cell range referenced
    #    in any Invoice/Invoice Log formula must either be prefixed with a
    #    sheet name, or (unprefixed) resolve to non-empty cells on its own
    #    sheet -- mirroring build/audit_products.py's bug detector for the
    #    "range silently resolves to the wrong sheet" class of bug.
    unprefixed_bad = []
    for ws in (invoice, log):
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                body = _strip_strings(v)
                for m in re.finditer(r"([A-Za-z0-9_' ]*!)?(\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+)", body):
                    prefix, rng = m.groups()
                    if prefix:
                        continue
                    try:
                        cells = ws[rng.replace("$", "")]
                    except Exception:
                        continue
                    flat = [x for rowc in cells for x in (rowc if isinstance(rowc, tuple) else (rowc,))]
                    filled = sum(1 for x in flat if x.value not in (None, ""))
                    if len(flat) > 3 and filled == 0:
                        unprefixed_bad.append((ws.title, c.coordinate, rng, v))
    assert not unprefixed_bad, f"unprefixed range(s) that resolve to an empty own-sheet range: {unprefixed_bad}"
    print("PASS: no unprefixed cross-sheet range bugs on Invoice or Invoice Log")

    # 5. Whitelisted functions only -- both common.py's broad list and the
    #    task brief's narrower explicit list.
    func_counter, disallowed_common, placeholder_hits = common.scan_workbook_formulas(wb)
    assert not disallowed_common, f"functions outside common.py's ALLOWED_FUNCTIONS: {disallowed_common}"
    disallowed_task = set(func_counter) - TASK_ALLOWED_FUNCTIONS
    assert not disallowed_task, f"functions outside the task brief's whitelist: {disallowed_task}"
    assert not placeholder_hits, f"leftover [placeholder] text found: {placeholder_hits}"
    print("PASS: functions used are all whitelisted:", sorted(func_counter))
    print("PASS: no leftover [placeholder] text")

    # 6. Print area set on Invoice, fits one page portrait.
    assert invoice.print_area, "Invoice sheet has no print_area set"
    assert invoice.page_setup.orientation == "portrait", "Invoice page orientation is not portrait"
    assert invoice.page_setup.fitToWidth == 1 and invoice.page_setup.fitToHeight == 1, (
        "Invoice page setup is not fit-to-one-page"
    )
    assert invoice.sheet_properties.pageSetUpPr.fitToPage, "Invoice sheet_properties.pageSetUpPr.fitToPage not set"
    print(f"PASS: Invoice print_area={invoice.print_area!r}, portrait, fit to one page")

    # 7. Sanity: 12 pre-formatted line-item rows (4 sample-filled), 30
    #    pre-formatted log rows (3 sample-filled).
    filled_items = sum(
        1 for r in range(ITEM_FIRST_ROW, ITEM_LAST_ROW + 1)
        if invoice.cell(row=r, column=1).value is not None
    )
    assert filled_items == ITEM_SAMPLE_ROWS, f"expected {ITEM_SAMPLE_ROWS} filled sample line items, found {filled_items}"
    assert (ITEM_LAST_ROW - ITEM_FIRST_ROW + 1) == ITEM_TOTAL_ROWS
    print(f"PASS: {ITEM_TOTAL_ROWS} line-item rows, {filled_items} sample rows filled")

    filled_log = sum(
        1 for r in range(LOG_FIRST_ROW, LOG_LAST_ROW + 1)
        if log.cell(row=r, column=1).value is not None
    )
    assert filled_log == LOG_SAMPLE_ROWS, f"expected {LOG_SAMPLE_ROWS} filled sample log rows, found {filled_log}"
    assert (LOG_LAST_ROW - LOG_FIRST_ROW + 1) == LOG_TOTAL_ROWS
    print(f"PASS: {LOG_TOTAL_ROWS} Invoice Log rows, {filled_log} sample rows filled")

    print("\nAll --verify checks passed.")


if __name__ == "__main__":
    if "--verify" in sys.argv:
        verify()
    else:
        build()
