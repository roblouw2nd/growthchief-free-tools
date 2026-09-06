#!/usr/bin/env python3
"""
Build the FREE "Simple Wedding Budget Spreadsheet" lead-magnet workbook.

This is a deliberately smaller SUBSET of the paid Wedding Budget & Planning
Toolkit (dist/p3/Wedding-Budget-Tracker.xlsx + Wedding-Planning-Checklist.xlsx
+ Wedding-Guest-List-Tracker.xlsx): 3 sheets instead of 3 separate workbooks
(11 sheets total across the toolkit), 12 fixed budget categories instead of
14 editable ones, a flat Payments log instead of a Vendor tracker with
quotes/deposits/balance due, and no Dashboard/charts, Timeline, Guest List,
or Seating tabs at all. It exists to be downloaded for free (no signup) and
to earn links/Pinterest saves, and to funnel interested users to the paid
Wedding Budget & Planning Toolkit and the custom-spreadsheet service.

Run with:    .venv/bin/python build/free_wedding.py
Verify with: .venv/bin/python build/free_wedding.py --verify

Follows SPEC.md: design system, Google Sheets compatibility rules, and the
shared Start Here sheet convention (via build/common.py).
"""
import re
import sys
from datetime import date

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT_PATH = ROOT / "dist" / "free" / "GrowthChief-Simple-Wedding-Budget.xlsx"

PRODUCT_NAME = "Simple Wedding Budget Spreadsheet (Free)"

ETSY_TOOLKIT_URL = "https://growthchief.etsy.com/listing/4560494732"
ETSY_CUSTOM_URL = "https://growthchief.etsy.com/listing/4564738495"

# ---------------------------------------------------------------------------
# Fixed vocab for the free version (deliberately not editable, unlike the
# paid toolkit's editable 14-category list on its Start Here tab)
# ---------------------------------------------------------------------------
CATEGORIES = [
    ("Venue", 6000), ("Catering", 4000), ("Photography", 2000), ("Attire", 1200),
    ("Flowers", 600), ("Music", 700), ("Stationery", 250), ("Rings", 1200),
    ("Transport", 400), ("Cake", 300), ("Hair & Makeup", 350), ("Other", 1000),
]
DEFAULT_TOTAL_BUDGET = 20000  # sum of CATEGORIES planned amounts is 18000,
                               # leaving a 2000 Unallocated sample tile

BUD_HEADER_ROW = 6
BUD_FIRST_ROW = 7
BUD_LAST_ROW = BUD_FIRST_ROW + len(CATEGORIES) - 1  # 18
BUD_TOTAL_ROW = BUD_LAST_ROW + 1  # 19
BUDGET_CATEGORY_RANGE = f"'Budget'!$A${BUD_FIRST_ROW}:$A${BUD_LAST_ROW}"

PAY_HEADER_ROW = 4
PAY_FIRST_ROW = 5
PAY_TOTAL_ROWS = 60
PAY_LAST_ROW = PAY_FIRST_ROW + PAY_TOTAL_ROWS - 1  # 64

# 6 sample payments, dated before "today" so they read as deposits already
# made against a wedding roughly nine months out.
SAMPLE_PAYMENTS = [
    (date(2026, 6, 1), "Venue", "Grand Oak Estate", 1500.0, "Deposit (example, replace me)"),
    (date(2026, 6, 15), "Photography", "Sunlight Studios", 500.0, "Deposit (example, replace me)"),
    (date(2026, 7, 1), "Rings", "Local Jewellers", 1200.0, "Paid in full (example, replace me)"),
    (date(2026, 7, 20), "Catering", "Table & Thyme Catering", 1000.0, "Deposit (example, replace me)"),
    (date(2026, 8, 5), "Stationery", "Paper & Ink Co", 250.0, "Paid in full (example, replace me)"),
    (date(2026, 8, 20), "Music", "DJ Marcus", 200.0, "Deposit (example, replace me)"),
]

link_font = Font(name=common.FONT_NAME, size=11, bold=True, color=common.ACCENT_COLOR, underline="single")
section_font = Font(name=common.FONT_NAME, size=13, bold=True, color=common.HEADER_FILL_COLOR)
TILE_VALUE_FONT = Font(name=common.FONT_NAME, size=20, bold=True, color=common.HEADER_FILL_COLOR)
TILE_LABEL_FONT = Font(name=common.FONT_NAME, size=11, bold=True, color="6B7280")


# ---------------------------------------------------------------------------
# Start Here
# ---------------------------------------------------------------------------

def build_start_here_sheet(wb):
    steps = [
        "Set your currency symbol, total budget, and wedding date in Settings below -- every sheet updates automatically.",
        "Go to the Budget tab and check the 12 categories -- adjust the Planned amount for each one to match your own wedding.",
        "Log every deposit and payment on the Payments tab as you make it: pick a Category from the dropdown, add the vendor, and enter the amount.",
        "Check the Budget tab any time -- Paid so far, Remaining, and % of Total update themselves, and any category you've overpaid turns red.",
        "Works in Excel and Google Sheets as-is -- no macros, no add-ons, nothing to enable.",
    ]
    settings = [
        {
            "label": "Currency symbol",
            "value": "$",
            "name": "CurrencySymbol",
            "note": "Used in every money header throughout this workbook. Doesn't convert currencies -- just changes the symbol shown.",
        },
        {
            "label": "Total wedding budget",
            "value": DEFAULT_TOTAL_BUDGET,
            "name": "TotalBudget",
            "number_format": common.NUMBER_FMT,
            "note": "Your overall target -- used for % of Total and the Unallocated tile on the Budget tab.",
        },
        {
            "label": "Wedding date",
            "value": date(2027, 6, 6),
            "name": "WeddingDate",
            "number_format": common.DATE_FMT,
            "note": "Used to work out Days to go, below.",
        },
    ]
    result = common.build_start_here(
        wb,
        PRODUCT_NAME,
        steps,
        settings,
        guide_link_placeholder=(
            "This is the free version -- it's simple by design, so there's no separate guide. "
            "If anything's unclear, the download page has a full walkthrough."
        ),
    )
    ws = result["ws"]
    row = result["settings_range"][1] + 2

    # --- Days to go tile ---
    label_cell = ws.cell(
        row=row, column=2,
        value='="Days to go until " & TEXT(WeddingDate,"mmm d, yyyy")',
    )
    label_cell.font = TILE_LABEL_FONT
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    row += 1

    value_cell = ws.cell(row=row, column=2, value="=WeddingDate-TODAY()")
    value_cell.font = TILE_VALUE_FONT
    value_cell.fill = common.accent_light_fill
    value_cell.border = common.thin_border
    value_cell.number_format = "0"
    value_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 30
    row += 2

    # --- Want more? ---
    header = ws.cell(row=row, column=2, value="Want more?")
    header.font = section_font
    row += 1

    intro = ws.cell(
        row=row, column=2,
        value=(
            "This free spreadsheet covers the basics: plan a budget, log payments. "
            "The full Wedding Budget & Planning Toolkit builds on top of it with:"
        ),
    )
    intro.font = common.body_font
    intro.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 32
    row += 1

    bullets = [
        "A bigger Budget & Expenses workbook with 14 editable categories, a Vendor field and Paid status "
        "on every expense line, and a Dashboard tab with a spend-by-category chart and a budget-vs-actual chart.",
        "A dedicated vendor tracker with quotes, deposits paid, and balance due worked out automatically for "
        "every vendor you're considering or have booked.",
        "A 60-task, 7-phase wedding planning timeline (12+ months before through wedding week) that tracks "
        "your progress automatically as you mark tasks done.",
        "A guest list tracker with RSVP, meal choice, and gift tracking, plus a seating chart tool that flags "
        "any table over capacity.",
    ]
    for b in bullets:
        c = ws.cell(row=row, column=2, value=f"-  {b}")
        c.font = common.body_font
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 30
        row += 1

    row += 1
    link1 = ws.cell(row=row, column=2, value="See the full Wedding Budget & Planning Toolkit on Etsy")
    link1.font = link_font
    link1.hyperlink = ETSY_TOOLKIT_URL
    row += 1
    link2 = ws.cell(row=row, column=2, value="Need something custom built instead? Custom spreadsheets, made to order")
    link2.font = link_font
    link2.hyperlink = ETSY_CUSTOM_URL
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    row += 2

    return ws


# ---------------------------------------------------------------------------
# Budget
# ---------------------------------------------------------------------------

def build_budget_sheet(wb):
    ws = common.new_data_sheet(wb, "Budget", tab_color=common.TAB_SLATE)
    common.set_col_widths(ws, {"A": 22, "B": 16, "C": 16, "D": 16, "E": 12})

    common.style_title(ws, "A1", "Budget")
    sub = ws.cell(row=2, column=1, value=(
        "Set what you plan to spend per category -- Paid so far updates automatically from the Payments tab."
    ))
    sub.font = common.subtitle_font
    sub.alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)
    ws.row_dimensions[2].height = 18

    # --- Unallocated tile ---
    label_cell = ws.cell(row=3, column=1, value='="Unallocated (" & CurrencySymbol & ")"')
    label_cell.font = TILE_LABEL_FONT
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=2)

    unallocated_cell = ws.cell(row=4, column=1, value=f"=TotalBudget-B{BUD_TOTAL_ROW}")
    unallocated_cell.font = TILE_VALUE_FONT
    unallocated_cell.fill = common.accent_light_fill
    unallocated_cell.border = common.thin_border
    unallocated_cell.number_format = common.NUMBER_FMT_NEG_RED
    unallocated_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=2)
    ws.row_dimensions[4].height = 30
    common.apply_negative_red(ws, "A4:B4")

    # --- Table ---
    common.style_header_row(ws, BUD_HEADER_ROW, 1, 5)
    headers = [
        "Category",
        '="Planned (" & CurrencySymbol & ")"',
        '="Paid so far (" & CurrencySymbol & ")"',
        '="Remaining (" & CurrencySymbol & ")"',
        "% of Total",
    ]
    for col, val in enumerate(headers, start=1):
        ws.cell(row=BUD_HEADER_ROW, column=col, value=val)

    pay_cat_range = f"Payments!$B${PAY_FIRST_ROW}:$B${PAY_LAST_ROW}"
    pay_amt_range = f"Payments!$D${PAY_FIRST_ROW}:$D${PAY_LAST_ROW}"

    for i, (cat, planned) in enumerate(CATEGORIES):
        r = BUD_FIRST_ROW + i
        ws.cell(row=r, column=1, value=cat)
        ws.cell(row=r, column=2, value=planned)
        ws.cell(row=r, column=3, value=f"=SUMIF({pay_cat_range},$A{r},{pay_amt_range})")
        ws.cell(row=r, column=4, value=f"=B{r}-C{r}")
        ws.cell(row=r, column=5, value=f"=IFERROR(B{r}/TotalBudget,0)")
        ws.cell(row=r, column=2).number_format = common.NUMBER_FMT
        ws.cell(row=r, column=3).number_format = common.NUMBER_FMT
        ws.cell(row=r, column=4).number_format = common.NUMBER_FMT
        ws.cell(row=r, column=5).number_format = common.PERCENT_FMT

    common.band_rows(ws, BUD_FIRST_ROW, BUD_LAST_ROW, 1, 5)

    ws.cell(row=BUD_TOTAL_ROW, column=1, value="Total")
    ws.cell(row=BUD_TOTAL_ROW, column=2, value=f"=SUM(B{BUD_FIRST_ROW}:B{BUD_LAST_ROW})")
    ws.cell(row=BUD_TOTAL_ROW, column=3, value=f"=SUM(C{BUD_FIRST_ROW}:C{BUD_LAST_ROW})")
    ws.cell(row=BUD_TOTAL_ROW, column=4, value=f"=B{BUD_TOTAL_ROW}-C{BUD_TOTAL_ROW}")
    ws.cell(row=BUD_TOTAL_ROW, column=5, value=f"=IFERROR(B{BUD_TOTAL_ROW}/TotalBudget,0)")
    for c in range(1, 6):
        cell = ws.cell(row=BUD_TOTAL_ROW, column=c)
        cell.fill = common.accent_fill
        cell.font = Font(name=common.FONT_NAME, size=11, bold=True, color=common.WHITE)
        cell.border = common.thin_border
    ws.cell(row=BUD_TOTAL_ROW, column=2).number_format = common.NUMBER_FMT
    ws.cell(row=BUD_TOTAL_ROW, column=3).number_format = common.NUMBER_FMT
    ws.cell(row=BUD_TOTAL_ROW, column=4).number_format = common.NUMBER_FMT
    ws.cell(row=BUD_TOTAL_ROW, column=5).number_format = common.PERCENT_FMT
    ws.row_dimensions[BUD_TOTAL_ROW].height = 22

    # Red fill on any category row where Paid so far > Planned (Remaining < 0)
    red_light = PatternFill("solid", fgColor="FEE2E2")
    ws.conditional_formatting.add(
        f"A{BUD_FIRST_ROW}:E{BUD_LAST_ROW}",
        FormulaRule(formula=[f"$D{BUD_FIRST_ROW}<0"], fill=red_light),
    )

    ws.freeze_panes = f"A{BUD_FIRST_ROW}"
    ws.sheet_view.zoomScale = 100
    return ws


# ---------------------------------------------------------------------------
# Payments
# ---------------------------------------------------------------------------

def build_payments_sheet(wb):
    ws = common.new_data_sheet(wb, "Payments", tab_color=common.TAB_SLATE)
    common.set_col_widths(ws, {"A": 13, "B": 16, "C": 22, "D": 14, "E": 34})

    common.style_title(ws, "A1", "Payments")
    sub = ws.cell(row=2, column=1, value=(
        "Log every deposit and payment here. Pick a Category, add the vendor, "
        "and enter the amount -- the Budget tab totals it automatically."
    ))
    sub.font = common.subtitle_font
    sub.alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)
    ws.row_dimensions[2].height = 18

    common.style_header_row(ws, PAY_HEADER_ROW, 1, 5)
    headers = ["Date", "Category", "Vendor", '="Amount (" & CurrencySymbol & ")"', "Note"]
    for col, val in enumerate(headers, start=1):
        ws.cell(row=PAY_HEADER_ROW, column=col, value=val)

    for i, (d, cat, vendor, amt, note) in enumerate(SAMPLE_PAYMENTS):
        r = PAY_FIRST_ROW + i
        ws.cell(row=r, column=1, value=d)
        ws.cell(row=r, column=2, value=cat)
        ws.cell(row=r, column=3, value=vendor)
        ws.cell(row=r, column=4, value=amt)
        ws.cell(row=r, column=5, value=note)

    common.band_rows(ws, PAY_FIRST_ROW, PAY_LAST_ROW, 1, 5)

    for r in range(PAY_FIRST_ROW, PAY_LAST_ROW + 1):
        ws.cell(row=r, column=1).number_format = common.DATE_FMT
        ws.cell(row=r, column=4).number_format = common.NUMBER_FMT_NEG_RED

    category_range = f"B{PAY_FIRST_ROW}:B{PAY_LAST_ROW}"
    amount_range = f"D{PAY_FIRST_ROW}:D{PAY_LAST_ROW}"

    # Point the dropdown at the live Budget category column via the
    # BudgetCategories defined name, rather than a duplicated literal list --
    # so the free workbook has exactly one place its 12 categories are typed in.
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1="BudgetCategories", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(category_range)

    common.apply_negative_red(ws, amount_range)

    common.freeze_header(ws, f"A{PAY_FIRST_ROW}")
    ws.sheet_view.zoomScale = 100
    return ws


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------

def build():
    wb = Workbook()
    build_start_here_sheet(wb)
    build_budget_sheet(wb)
    build_payments_sheet(wb)

    assert wb.sheetnames == ["Start Here", "Budget", "Payments"], wb.sheetnames

    wb.defined_names["BudgetCategories"] = DefinedName(
        "BudgetCategories", attr_text=BUDGET_CATEGORY_RANGE)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} ({OUT_PATH.stat().st_size:,} bytes)")
    return OUT_PATH


# ---------------------------------------------------------------------------
# Verify
# ---------------------------------------------------------------------------

TASK_ALLOWED_FUNCTIONS = {
    "SUM", "SUMIF", "SUMIFS", "COUNTIF", "IF", "IFERROR", "TODAY", "DATE",
    "TEXT", "EOMONTH", "ROUND",
}

_RANGE_RE = re.compile(r"(?<![!\w$'])(\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+)")


def _strip_strings(formula):
    return re.sub(r'"[^"]*"', "", formula)


def verify():
    assert OUT_PATH.exists(), f"{OUT_PATH} does not exist -- run without --verify first"
    wb = load_workbook(OUT_PATH)

    # 1. Sheet names, exactly 3, in order.
    assert wb.sheetnames == ["Start Here", "Budget", "Payments"], (
        f"unexpected sheet list: {wb.sheetnames}"
    )
    print("PASS: sheet names ==", wb.sheetnames)

    # 2. Defined names.
    for name in ("CurrencySymbol", "TotalBudget", "WeddingDate", "BudgetCategories"):
        assert name in wb.defined_names, f"Missing defined name: {name}"
    print("PASS: defined names present: CurrencySymbol, TotalBudget, WeddingDate, BudgetCategories")

    start = wb["Start Here"]
    days_formula = None
    for row in start.iter_rows():
        for c in row:
            if c.value == "=WeddingDate-TODAY()":
                days_formula = c.coordinate
    assert days_formula, "Start Here missing 'Days to go' formula =WeddingDate-TODAY()"
    print(f"PASS: Days to go formula found at Start Here!{days_formula}")

    budget = wb["Budget"]
    payments = wb["Payments"]

    # 3. Budget table: 12 categories, formulas wired correctly.
    cats = [budget.cell(row=r, column=1).value for r in range(BUD_FIRST_ROW, BUD_LAST_ROW + 1)]
    assert cats == [c for c, _ in CATEGORIES], f"category mismatch: {cats}"
    assert len(cats) == 12, f"expected 12 categories, got {len(cats)}"
    print("PASS: 12 Budget categories match spec:", cats)

    assert budget.cell(row=BUD_FIRST_ROW, column=3).value == (
        f"=SUMIF(Payments!$B${PAY_FIRST_ROW}:$B${PAY_LAST_ROW},$A{BUD_FIRST_ROW},"
        f"Payments!$D${PAY_FIRST_ROW}:$D${PAY_LAST_ROW})"
    ), f"unexpected Paid-so-far formula: {budget.cell(row=BUD_FIRST_ROW, column=3).value!r}"
    assert budget.cell(row=BUD_FIRST_ROW, column=4).value == f"=B{BUD_FIRST_ROW}-C{BUD_FIRST_ROW}"
    assert budget.cell(row=BUD_FIRST_ROW, column=5).value == f"=IFERROR(B{BUD_FIRST_ROW}/TotalBudget,0)"
    print("PASS: Budget row formulas (Paid so far / Remaining / % of Total) correct")

    assert budget["A4"].value == f"=TotalBudget-B{BUD_TOTAL_ROW}", (
        f"Unallocated tile formula wrong: {budget['A4'].value!r}"
    )
    print("PASS: Unallocated tile formula correct")

    assert budget.cell(row=BUD_TOTAL_ROW, column=2).value == f"=SUM(B{BUD_FIRST_ROW}:B{BUD_LAST_ROW})"
    assert budget.cell(row=BUD_TOTAL_ROW, column=3).value == f"=SUM(C{BUD_FIRST_ROW}:C{BUD_LAST_ROW})"
    print("PASS: Budget totals row formulas correct")

    budget_cf_count = sum(len(rules) for rules in budget.conditional_formatting._cf_rules.values())
    assert budget_cf_count == 2, f"expected 2 conditional-format rule groups on Budget (Unallocated red + over-budget red), got {budget_cf_count}"
    print(f"PASS: {budget_cf_count} conditional-format rule group(s) present on Budget")

    default_sum = sum(v for _, v in CATEGORIES)
    assert default_sum == 18000, f"sample planned amounts should sum to 18000, got {default_sum}"
    print("PASS: sample Planned amounts sum to", default_sum)

    # 4. Payments sheet: dropdown, 60 pre-formatted rows, 6 sample rows.
    dvs = payments.data_validations.dataValidation
    assert len(dvs) == 1, f"expected 1 data validation on Payments, found {len(dvs)}"
    dv = dvs[0]
    assert dv.type == "list", f"unexpected validation type {dv.type!r}"
    assert dv.formula1 == "BudgetCategories", f"Category dropdown should reference BudgetCategories, got {dv.formula1!r}"
    assert not dv.formula1.startswith("="), "DataValidation formula1 must not lead with '='"
    assert f"B{PAY_FIRST_ROW}:B{PAY_LAST_ROW}" in str(dv.sqref), f"Category dropdown range missing, sqref={dv.sqref}"
    print("PASS: Payments Category dropdown wired to BudgetCategories:", dv.sqref)

    filled = sum(
        1 for r in range(PAY_FIRST_ROW, PAY_LAST_ROW + 1)
        if payments.cell(row=r, column=1).value is not None
    )
    assert filled == len(SAMPLE_PAYMENTS), f"expected {len(SAMPLE_PAYMENTS)} filled sample rows, found {filled}"
    total_rows = PAY_LAST_ROW - PAY_FIRST_ROW + 1
    assert total_rows == PAY_TOTAL_ROWS == 60, total_rows
    print(f"PASS: {PAY_TOTAL_ROWS} pre-formatted Payments rows, {filled} sample rows filled")

    # 5. No unprefixed cross-sheet ranges anywhere in the workbook: every
    #    unprefixed multi-cell range referenced in a formula must actually be
    #    a same-sheet range (its row numbers must fall inside that sheet's own
    #    data block), and every formula that legitimately reaches onto another
    #    sheet (the Budget SUMIF over Payments) must carry an explicit prefix.
    same_sheet_ranges = {
        "Budget": (BUD_FIRST_ROW, BUD_LAST_ROW),
        "Payments": (PAY_FIRST_ROW, PAY_LAST_ROW),
    }
    bad_unprefixed = []
    for ws in wb.worksheets:
        lo_hi = same_sheet_ranges.get(ws.title)
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                body = _strip_strings(v)
                for m in re.finditer(r"([A-Za-z0-9_']*!)?(\$?[A-Z]{1,3}\$?\d+):(\$?[A-Z]{1,3}\$?\d+)", body):
                    prefix, ref1, ref2 = m.groups()
                    if prefix:
                        continue  # explicitly sheet-prefixed -- fine
                    r1 = int(re.sub(r"\D", "", ref1))
                    r2 = int(re.sub(r"\D", "", ref2))
                    if lo_hi and lo_hi[0] <= r1 and r2 <= lo_hi[1]:
                        continue  # genuinely a same-sheet range (e.g. a SUM total)
                    bad_unprefixed.append((ws.title, c.coordinate, f"{ref1}:{ref2}", v))
    assert not bad_unprefixed, f"unprefixed cross-sheet range(s) found: {bad_unprefixed}"
    print("PASS: no unprefixed cross-sheet ranges anywhere in the workbook")

    payments_refs = sum(
        1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=") and "Payments!" in c.value
    )
    assert payments_refs == len(CATEGORIES), (
        f"expected {len(CATEGORIES)} sheet-prefixed 'Payments!' SUMIF formulas on Budget, found {payments_refs}"
    )
    print(f"PASS: {payments_refs} sheet-prefixed Payments! references on Budget (one per category)")

    # 6. Whitelisted functions only.
    func_counter, disallowed_common, placeholder_hits = common.scan_workbook_formulas(wb)
    assert not disallowed_common, f"functions outside common.py's ALLOWED_FUNCTIONS: {disallowed_common}"
    disallowed_task = set(func_counter) - TASK_ALLOWED_FUNCTIONS
    assert not disallowed_task, f"functions outside the task brief's whitelist: {disallowed_task}"
    assert not placeholder_hits, f"leftover [placeholder] text found: {placeholder_hits}"
    print("PASS: functions used are all whitelisted:", sorted(func_counter))
    print("PASS: no leftover [placeholder] text")

    print("\nAll --verify checks passed.")


if __name__ == "__main__":
    if "--verify" in sys.argv:
        verify()
    else:
        build()
