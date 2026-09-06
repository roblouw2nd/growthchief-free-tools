#!/usr/bin/env python3
"""
Build the FREE "Simple Etsy Fee & Profit Sheet" lead-magnet workbook.

Deliberately smaller SUBSET of the paid Etsy Seller Bookkeeping Toolkit
(dist/p2/Etsy-Fee-Profit-Calculator.xlsx, dist/p2/Etsy-Seller-Bookkeeping-
Tracker.xlsx, dist/p2/Etsy-Order-Inventory-Tracker.xlsx): a single-item fee
& profit calculator plus a flat 100-row Sales Log instead of a 10-item
Product Comparison sheet, a 12-month Bookkeeping Tracker with Dashboard, and
an Order & Inventory Tracker with low-stock alerts. It exists to be
downloaded for free (no signup) and to earn links/Pinterest saves, and to
funnel interested Etsy sellers to the paid toolkit.

Uses the EXACT SAME fee model and default values as the paid Etsy Fee &
Profit Calculator (dist/p2/Etsy-Fee-Profit-Calculator.xlsx, built by
build/p2_fee_profit_calculator.py): a $0.20 flat listing fee, 6.5%
transaction fee, 3% + $0.25 payment processing fee. Every rate is an
editable Settings input on Start Here -- nothing is hardcoded into a
formula -- plus one addition the paid calculator's Calculator sheet doesn't
model: an Offsite Ads fee %, defaulted to 0 and applied only when a given
sale is flagged as an Offsite Ads sale, since (per marketing/site/
etsy-fee-calculator/index.html's own FAQ) that fee doesn't apply to every
sale the same way.

Run with:    .venv/bin/python build/free_etsy.py
Verify with: .venv/bin/python build/free_etsy.py --verify

Follows SPEC.md: design system, Google Sheets compatibility rules, and the
shared Start Here sheet convention (via build/common.py).
"""
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT_PATH = ROOT / "dist" / "free" / "GrowthChief-Simple-Etsy-Profit-Sheet.xlsx"

PRODUCT_NAME = "Simple Etsy Fee & Profit Sheet (Free)"

ETSY_P2_URL = "https://growthchief.etsy.com/listing/4560467662"
ETSY_CUSTOM_URL = "https://growthchief.etsy.com/listing/4564738495"

link_font = Font(name=common.FONT_NAME, size=11, bold=True, color=common.ACCENT_COLOR, underline="single")
section_font = Font(name=common.FONT_NAME, size=13, bold=True, color=common.HEADER_FILL_COLOR)

# QA fix precedent from build/p2_fee_profit_calculator.py (2026-08-22):
# PERCENT_FMT ("0%", zero decimals) rounds a 6.5% fee to a displayed "7%" --
# misleading on a fee-accuracy product even though the underlying float
# (0.065) stays exact. One decimal place for every fee/margin percentage.
PERCENT_FMT_1DP = "0.0%"

FEE_NOTE = ("Check Etsy's current fee schedule -- fees change. This is your own "
            "editable estimate, not official Etsy data.")


def money_header(label):
    return f'="{label} (" & CurrencySymbol & ")"'


# ---------------------------------------------------------------------------
# Start Here
# ---------------------------------------------------------------------------

def build_start_here_sheet(wb):
    steps = [
        "Set your currency symbol and Etsy fee settings below -- these are your own editable "
        "estimates, so check Etsy's current fee schedule and update them any time fees change.",
        "Open the Profit Calculator tab and fill in one sale: sale price, shipping you charge "
        "the buyer, item cost, packaging cost, shipping you actually pay, and whether the sale "
        "came through Etsy Offsite Ads.",
        "The sheet works out your total Etsy fees and true net profit -- and the reverse line "
        "tells you what to charge to hit a target margin.",
        "Log every sale on the Sales Log tab as it happens -- Fees and Net fill themselves in "
        "from the same fee settings.",
        "Works in Excel and Google Sheets as-is -- no macros, no add-ons, nothing to enable.",
    ]
    settings = [
        {
            "label": "Currency symbol", "value": "$", "name": "CurrencySymbol",
            "note": "Display only. Number formats stay plain so this works in any currency.",
        },
        {
            "label": "Listing fee (flat, per listing)", "value": 0.20, "name": "ListingFee",
            "number_format": common.NUMBER_FMT, "note": FEE_NOTE,
        },
        {
            "label": "Transaction fee %", "value": 0.065, "name": "TransactionPct",
            "number_format": PERCENT_FMT_1DP, "note": FEE_NOTE,
        },
        {
            "label": "Payment processing fee %", "value": 0.03, "name": "ProcessingPct",
            "number_format": PERCENT_FMT_1DP, "note": FEE_NOTE,
        },
        {
            "label": "Payment processing flat fee", "value": 0.25, "name": "ProcessingFixed",
            "number_format": common.NUMBER_FMT, "note": FEE_NOTE,
        },
        {
            "label": "Offsite Ads fee %", "value": 0.0, "name": "OffsiteAdsPct",
            "number_format": PERCENT_FMT_1DP,
            "note": ("0% by default -- only applies when you flag a sale as an Offsite Ads sale. "
                     "Set to 0.12 or 0.15 (12% or 15%) if you're opted in; check Etsy's current "
                     "threshold and rate, since this changes."),
        },
    ]
    result = common.build_start_here(
        wb,
        PRODUCT_NAME,
        steps,
        settings,
        toolkit_name="Etsy Seller Bookkeeping Toolkit",
        guide_link_placeholder=(
            "This is the free version -- it's simple by design, so there's no separate guide. "
            "If anything's unclear, the download page has a full walkthrough."
        ),
        extra_note=(
            "This tool does plain arithmetic based on numbers you provide. It isn't tax, legal, "
            "or business advice, and the fee defaults above are estimates, not official Etsy "
            "figures -- always check Etsy's current fee schedule."
        ),
    )
    ws = result["ws"]
    ws.sheet_properties.tabColor = common.TAB_SLATE
    row = ws.max_row + 2

    header = ws.cell(row=row, column=2, value="Want more?")
    header.font = section_font
    row += 1

    intro = ws.cell(
        row=row, column=2,
        value=(
            "This free sheet covers one sale at a time plus a flat sales log. The full Etsy "
            "Seller Bookkeeping Toolkit builds on the same fee model with:"
        ),
    )
    intro.font = common.body_font
    intro.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 32
    row += 1

    bullets = [
        "This same fee & profit calculator, plus a 10-product Product Comparison sheet to run "
        "the same numbers across up to 10 product ideas side by side.",
        "A full Bookkeeping Tracker with a year-at-a-glance Dashboard and 12 monthly tabs "
        "(Jan-Dec), instead of one flat log.",
        "An Order & Inventory Tracker with its own Dashboard, an Orders log, and Inventory "
        "tracking with automatic low-stock reorder flags.",
    ]
    for b in bullets:
        c = ws.cell(row=row, column=2, value=f"-  {b}")
        c.font = common.body_font
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 32
        row += 1

    row += 1
    link1 = ws.cell(row=row, column=2, value="See the full Etsy Seller Bookkeeping Toolkit on Etsy")
    link1.font = link_font
    link1.hyperlink = ETSY_P2_URL
    row += 1
    link2 = ws.cell(row=row, column=2, value="Need something custom built for your shop? Custom spreadsheets, made to order")
    link2.font = link_font
    link2.hyperlink = ETSY_CUSTOM_URL
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    row += 2

    return ws


# ---------------------------------------------------------------------------
# Profit Calculator
# ---------------------------------------------------------------------------

def build_calculator_sheet(wb):
    ws = common.new_data_sheet(wb, "Profit Calculator", tab_color=common.TAB_TEAL, gridlines=False)
    common.set_col_widths(ws, {"A": 3, "B": 34, "C": 16, "D": 42, "E": 3})

    ws.merge_cells("B2:D2")
    common.style_title(ws, "B2", "Profit Calculator")
    ws.row_dimensions[2].height = 26
    ws.merge_cells("B3:D3")
    ws.cell(row=3, column=2,
            value="Price one Etsy sale, see what Etsy actually takes, and find out what to "
                  "charge to hit the margin you want.").font = common.note_font

    # --- Inputs -----------------------------------------------------------
    row = 5
    ws.cell(row=row, column=2, value="Your inputs").font = Font(
        name=common.FONT_NAME, size=13, bold=True, color=common.ACCENT_COLOR)
    row += 1
    common.style_header_row(ws, row, 2, 4, height=22)
    ws.cell(row=row, column=2, value="Input")
    ws.cell(row=row, column=3, value="Value")
    ws.cell(row=row, column=4, value="Notes")
    row += 1

    inputs = [
        (money_header("Sale price"), 28.00, common.NUMBER_FMT, "SalePrice",
         "What the buyer pays for the item itself, before shipping."),
        (money_header("Shipping charged to buyer"), 4.50, common.NUMBER_FMT, "ShippingCharged",
         "What you charge the buyer for shipping -- this counts toward Etsy's fees too."),
        (money_header("Item cost"), 9.00, common.NUMBER_FMT, "ItemCost",
         "What the item itself (materials, stock, etc.) cost you to make or buy in."),
        (money_header("Packaging cost"), 1.20, common.NUMBER_FMT, "PackagingCost",
         "Boxes, tissue, labels, thank-you cards -- whatever you use to ship it."),
        (money_header("Shipping cost paid"), 4.50, common.NUMBER_FMT, "ShippingCostPaid",
         "What the shipping label actually costs you -- a real expense, and not the same "
         "number as shipping charged above. Both are needed to see your true net profit."),
        ("Offsite ad sale?", "No", None, "OffsiteAdSale",
         "Yes if this sale came through an Etsy Offsite Ads click (Etsy charges an extra fee "
         "on those only). No for a normal on-Etsy sale."),
        ("Target profit margin %", 0.30, PERCENT_FMT_1DP, "TargetMarginPct",
         "Used by the reverse line below to work out what to charge."),
    ]
    input_rows = {}
    offsite_dv_row = None
    for label, value, fmt, name, note in inputs:
        ws.cell(row=row, column=2, value=label).font = common.label_font
        val_cell = ws.cell(row=row, column=3, value=value)
        if fmt:
            val_cell.number_format = fmt
        val_cell.font = Font(name=common.FONT_NAME, size=11, bold=True, color=common.ACCENT_COLOR)
        val_cell.fill = common.accent_light_fill
        val_cell.alignment = Alignment(horizontal="center")
        note_cell = ws.cell(row=row, column=4, value=note)
        note_cell.font = common.note_font
        note_cell.alignment = Alignment(wrap_text=True, vertical="center")
        for c in range(2, 5):
            ws.cell(row=row, column=c).border = common.thin_border
        ws.row_dimensions[row].height = 30
        common.add_defined_name(wb, name, "Profit Calculator", f"C{row}")
        input_rows[name] = row
        if name == "OffsiteAdSale":
            offsite_dv_row = row
        row += 1
    row += 1

    common.add_list_validation(ws, f"C{offsite_dv_row}", ["Yes", "No"])

    # --- Fee & profit breakdown -------------------------------------------
    ws.cell(row=row, column=2, value="Fee & profit breakdown").font = Font(
        name=common.FONT_NAME, size=13, bold=True, color=common.ACCENT_COLOR)
    row += 1
    common.style_header_row(ws, row, 2, 4, height=22)
    ws.cell(row=row, column=2, value="Line")
    ws.cell(row=row, column=3, value="Amount")
    ws.cell(row=row, column=4, value="How it's calculated")
    row += 1

    lines = [
        (money_header("Revenue (sale + shipping)"), "=SalePrice+ShippingCharged", common.NUMBER_FMT,
         "Revenue", "Sale price + shipping charged"),
        (money_header("Etsy listing fee"), "=ListingFee", common.NUMBER_FMT,
         "ListingFeeAmt", "Flat fee, from Settings"),
        (money_header("Etsy transaction fee"), "=Revenue*TransactionPct", common.NUMBER_FMT,
         "TransactionFeeAmt", "Revenue x transaction fee %"),
        (money_header("Payment processing fee"), "=Revenue*ProcessingPct+ProcessingFixed", common.NUMBER_FMT,
         "ProcessingFeeAmt", "(Revenue x processing fee %) + processing flat fee"),
        (money_header("Offsite Ads fee"), '=IF(OffsiteAdSale="Yes",Revenue*OffsiteAdsPct,0)', common.NUMBER_FMT,
         "OffsiteFeeAmt", "Revenue x Offsite Ads fee %, only if this sale is flagged Yes"),
        (money_header("Total Etsy fees"), "=ListingFee+TransactionFeeAmt+ProcessingFeeAmt+OffsiteFeeAmt", common.NUMBER_FMT,
         "TotalFees", "Listing + transaction + processing + Offsite Ads fees"),
        (money_header("Net profit"), "=Revenue-ItemCost-PackagingCost-ShippingCostPaid-TotalFees", common.NUMBER_FMT,
         "NetProfit", "Revenue - item cost - packaging - shipping paid - total fees"),
        ("Profit margin %", '=IFERROR(NetProfit/Revenue,"Check inputs")', PERCENT_FMT_1DP,
         "MarginPct", "Net profit / revenue"),
    ]
    for label, formula, fmt, name, explanation in lines:
        ws.cell(row=row, column=2, value=label).font = common.label_font
        val_cell = ws.cell(row=row, column=3, value=formula)
        val_cell.number_format = fmt
        val_cell.font = Font(name=common.FONT_NAME, size=13, bold=True, color=common.ACCENT_COLOR)
        val_cell.fill = common.accent_light_fill
        val_cell.alignment = Alignment(horizontal="center")
        note_cell = ws.cell(row=row, column=4, value=explanation)
        note_cell.font = common.note_font
        note_cell.alignment = Alignment(wrap_text=True, vertical="center")
        for c in range(2, 5):
            ws.cell(row=row, column=c).border = common.thin_border
        ws.row_dimensions[row].height = 26
        common.add_defined_name(wb, name, "Profit Calculator", f"C{row}")
        row += 1
    row += 1

    # --- Reverse: price needed for target margin ---------------------------
    ws.cell(row=row, column=2, value="What to charge to hit your target margin").font = Font(
        name=common.FONT_NAME, size=13, bold=True, color=common.ACCENT_COLOR)
    row += 1
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    ws.cell(row=row, column=2,
            value="Keeps shipping charged the same as above and solves for the sale price "
                  "needed to hit your target profit margin %, after fees, item cost, packaging, "
                  "and shipping paid.").font = common.note_font
    row += 1
    common.style_header_row(ws, row, 2, 4, height=22)
    ws.cell(row=row, column=2, value="Result")
    ws.cell(row=row, column=3, value="Value")
    ws.cell(row=row, column=4, value="How it's calculated")
    row += 1

    reverse_lines = [
        (money_header("Revenue needed"),
         '=IFERROR((ListingFee+ProcessingFixed+ItemCost+PackagingCost+ShippingCostPaid)'
         '/(1-TransactionPct-ProcessingPct-IF(OffsiteAdSale="Yes",OffsiteAdsPct,0)-TargetMarginPct),'
         '"Check inputs")',
         common.NUMBER_FMT, "RevenueNeeded",
         "(Listing fee + processing flat fee + item cost + packaging + shipping paid) / "
         "(1 - transaction fee % - processing fee % - Offsite Ads fee % if flagged - target margin %)"),
        (money_header("Sale price needed"),
         '=IFERROR(RevenueNeeded-ShippingCharged,"Check inputs")',
         common.NUMBER_FMT, "SalePriceNeeded",
         "Revenue needed - shipping charged (shipping stays the same as your inputs above)"),
    ]
    for label, formula, fmt, name, explanation in reverse_lines:
        ws.cell(row=row, column=2, value=label).font = common.label_font
        val_cell = ws.cell(row=row, column=3, value=formula)
        val_cell.number_format = fmt
        val_cell.font = Font(name=common.FONT_NAME, size=13, bold=True, color=common.ACCENT_COLOR)
        val_cell.fill = common.accent_light_fill
        val_cell.alignment = Alignment(horizontal="center")
        note_cell = ws.cell(row=row, column=4, value=explanation)
        note_cell.font = common.note_font
        note_cell.alignment = Alignment(wrap_text=True, vertical="center")
        for c in range(2, 5):
            ws.cell(row=row, column=c).border = common.thin_border
        ws.row_dimensions[row].height = 30
        common.add_defined_name(wb, name, "Profit Calculator", f"C{row}")
        row += 1

    common.apply_negative_red(ws, f"C{input_rows['SalePrice']}:C{row - 1}")
    common.freeze_header(ws, "B6")
    ws.sheet_view.zoomScale = 100
    return ws


# ---------------------------------------------------------------------------
# Sales Log
# ---------------------------------------------------------------------------

SL_HEADER_ROW = 7
SL_FIRST_DATA_ROW = 8
SL_TOTAL_ROWS = 100
SL_LAST_DATA_ROW = SL_FIRST_DATA_ROW + SL_TOTAL_ROWS - 1  # 107
SL_TOTAL_ROW_NUM = SL_LAST_DATA_ROW + 2  # one blank row, then Totals -- 109

# 5 sample rows, dated in the current real month so the "This month's net"
# SUMIFS tile is non-zero when opened/recalculated today, not just an
# all-time total. (date, item, sale price, shipping charged, item cost,
# shipping paid, offsite ad? Yes/No)
SAMPLE_LOG_ROWS = [
    (date(2026, 9, 1), "Personalised Wooden Sign (example, replace me)", 28.00, 4.50, 9.00, 4.50, "No"),
    (date(2026, 9, 2), "Custom Pet Portrait Print (example, replace me)", 45.00, 6.00, 12.00, 5.50, "Yes"),
    (date(2026, 9, 3), "Macrame Wall Hanging (example, replace me)", 32.00, 5.00, 10.00, 4.75, "No"),
    (date(2026, 9, 4), "Engraved Cutting Board (example, replace me)", 38.00, 7.50, 14.00, 7.00, "Yes"),
    (date(2026, 9, 5), "Beaded Bracelet Set (example, replace me)", 22.00, 3.50, 6.00, 3.25, "No"),
]


def build_sales_log_sheet(wb):
    ws = common.new_data_sheet(wb, "Sales Log", tab_color=common.TAB_SLATE)
    common.set_col_widths(ws, {
        "A": 13, "B": 34, "C": 13, "D": 13, "E": 13, "F": 13, "G": 12, "H": 12, "I": 12,
    })

    common.style_title(ws, "A1", "Sales Log")
    sub = ws.cell(row=2, column=1, value=(
        "Log every sale here. Fees and Net fill themselves in from the fee settings on "
        "Start Here -- flag Offsite ad sales so the Offsite Ads fee applies only where it should."
    ))
    sub.font = common.subtitle_font
    sub.alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=9)
    ws.row_dimensions[2].height = 18

    date_range = f"A{SL_FIRST_DATA_ROW}:A{SL_LAST_DATA_ROW}"
    net_range = f"I{SL_FIRST_DATA_ROW}:I{SL_LAST_DATA_ROW}"
    month_start = "(EOMONTH(TODAY(),-1)+1)"
    month_end_exclusive = "(EOMONTH(TODAY(),0)+1)"

    tile_label = ws.cell(row=4, column=1,
                          value='="This month\'s net (" & TEXT(TODAY(),"mmmm yyyy") & ")"')
    tile_label.font = Font(name=common.FONT_NAME, size=11, bold=True, color="6B7280")
    tile_value = ws.cell(row=5, column=1, value=(
        f'=SUMIFS({net_range},{date_range},">="&{month_start},{date_range},"<"&{month_end_exclusive})'
    ))
    tile_value.font = Font(name=common.FONT_NAME, size=20, bold=True, color=common.HEADER_FILL_COLOR)
    tile_value.fill = common.accent_light_fill
    tile_value.border = common.thin_border
    tile_value.number_format = common.NUMBER_FMT_NEG_RED
    tile_value.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=3)
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=3)
    ws.row_dimensions[5].height = 30
    common.apply_negative_red(ws, "A5:C5")

    common.style_header_row(ws, SL_HEADER_ROW, 1, 9)
    headers = [
        "Date", "Item",
        money_header("Sale price"), money_header("Shipping charged"),
        money_header("Item cost"), money_header("Shipping paid"),
        "Offsite ad?", money_header("Fees"), money_header("Net"),
    ]
    for col, val in enumerate(headers, start=1):
        ws.cell(row=SL_HEADER_ROW, column=col, value=val)

    for i, (d, item, sale_price, ship_charged, item_cost, ship_paid, offsite) in enumerate(SAMPLE_LOG_ROWS):
        r = SL_FIRST_DATA_ROW + i
        ws.cell(row=r, column=1, value=d)
        ws.cell(row=r, column=2, value=item)
        ws.cell(row=r, column=3, value=sale_price)
        ws.cell(row=r, column=4, value=ship_charged)
        ws.cell(row=r, column=5, value=item_cost)
        ws.cell(row=r, column=6, value=ship_paid)
        ws.cell(row=r, column=7, value=offsite)

    for r in range(SL_FIRST_DATA_ROW, SL_LAST_DATA_ROW + 1):
        revenue = f"(C{r}+D{r})"
        fees_formula = (
            f'=IF($B{r}="","",IFERROR(ListingFee+{revenue}*TransactionPct+'
            f'({revenue}*ProcessingPct+ProcessingFixed)+'
            f'IF(G{r}="Yes",{revenue}*OffsiteAdsPct,0),"Check inputs"))'
        )
        net_formula = f'=IF($B{r}="","",IFERROR({revenue}-E{r}-F{r}-H{r},"Check inputs"))'
        ws.cell(row=r, column=8, value=fees_formula)
        ws.cell(row=r, column=9, value=net_formula)

    common.band_rows(ws, SL_FIRST_DATA_ROW, SL_LAST_DATA_ROW, 1, 9)

    for r in range(SL_FIRST_DATA_ROW, SL_LAST_DATA_ROW + 1):
        ws.cell(row=r, column=1).number_format = common.DATE_FMT
        for col in (3, 4, 5, 6, 8, 9):
            ws.cell(row=r, column=col).number_format = common.NUMBER_FMT_NEG_RED
        ws.cell(row=r, column=7).alignment = Alignment(horizontal="center")

    offsite_range = f"G{SL_FIRST_DATA_ROW}:G{SL_LAST_DATA_ROW}"
    common.add_list_validation(ws, offsite_range, ["Yes", "No"])

    # Totals row
    total_row = SL_TOTAL_ROW_NUM
    ws.cell(row=total_row, column=1, value="Total")
    for col in (3, 4, 5, 6, 8, 9):
        letter = ws.cell(row=SL_FIRST_DATA_ROW, column=col).coordinate[0]
        cell = ws.cell(row=total_row, column=col,
                        value=f"=SUM({letter}{SL_FIRST_DATA_ROW}:{letter}{SL_LAST_DATA_ROW})")
        cell.number_format = common.NUMBER_FMT_NEG_RED
    for col in range(1, 10):
        cell = ws.cell(row=total_row, column=col)
        cell.fill = common.accent_fill
        cell.font = Font(name=common.FONT_NAME, size=11, bold=True, color=common.WHITE)
        cell.border = common.thin_border
    ws.row_dimensions[total_row].height = 22

    common.freeze_header(ws, f"A{SL_FIRST_DATA_ROW}")
    ws.sheet_view.zoomScale = 100
    return ws, total_row


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------

def build():
    wb = Workbook()
    build_start_here_sheet(wb)
    build_calculator_sheet(wb)
    build_sales_log_sheet(wb)

    assert wb.sheetnames == ["Start Here", "Profit Calculator", "Sales Log"], wb.sheetnames

    wb.active = wb["Start Here"]
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} ({OUT_PATH.stat().st_size:,} bytes)")
    return OUT_PATH


# ---------------------------------------------------------------------------
# Verify
# ---------------------------------------------------------------------------

TASK_ALLOWED_FUNCTIONS = {
    "SUM", "SUMIF", "SUMIFS", "COUNTIF", "IF", "IFERROR", "TODAY", "DATE",
    "TEXT", "EOMONTH", "ROUND",
}


def verify():
    assert OUT_PATH.exists(), f"{OUT_PATH} does not exist -- run without --verify first"
    wb = load_workbook(OUT_PATH)

    # 1. Sheet names, exactly 3, in order.
    assert wb.sheetnames == ["Start Here", "Profit Calculator", "Sales Log"], (
        f"unexpected sheet list: {wb.sheetnames}"
    )
    print("PASS: sheet names ==", wb.sheetnames)

    calc = wb["Profit Calculator"]
    log = wb["Sales Log"]

    # 2. Defined names present.
    required_names = (
        "CurrencySymbol", "ListingFee", "TransactionPct", "ProcessingPct",
        "ProcessingFixed", "OffsiteAdsPct",
        "SalePrice", "ShippingCharged", "ItemCost", "PackagingCost",
        "ShippingCostPaid", "OffsiteAdSale", "TargetMarginPct",
        "Revenue", "TotalFees", "NetProfit", "MarginPct",
        "RevenueNeeded", "SalePriceNeeded",
    )
    for name in required_names:
        assert name in wb.defined_names, f"Missing defined name: {name}"
    print(f"PASS: all {len(required_names)} required defined names present")

    # 3. Settings sample values match the paid Etsy Fee & Profit Calculator's
    #    defaults (dist/p2/Etsy-Fee-Profit-Calculator.xlsx), so the free sheet
    #    never contradicts the paid one -- plus OffsiteAdsPct default 0.
    start = wb["Start Here"]

    def resolve(name):
        dn = wb.defined_names[name]
        ref = dn.attr_text.split("!")[-1].replace("$", "")
        sheet = dn.attr_text.split("!")[0].strip("'")
        return wb[sheet][ref]

    assert resolve("CurrencySymbol").value == "$"
    assert resolve("ListingFee").value == 0.20
    assert resolve("TransactionPct").value == 0.065
    assert resolve("ProcessingPct").value == 0.03
    assert resolve("ProcessingFixed").value == 0.25
    assert resolve("OffsiteAdsPct").value == 0.0
    print("PASS: Settings defaults match the paid Etsy Fee & Profit Calculator "
          "(ListingFee=0.20, TransactionPct=0.065, ProcessingPct=0.03, "
          "ProcessingFixed=0.25) plus OffsiteAdsPct default 0")

    # 4. No fee value hardcoded into the fee formulas -- must reference
    #    defined names (same discipline as build/p2_fee_profit_calculator.py).
    import re as _re
    fee_formula_cells = []
    for ws in (calc, log):
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and (
                    "TransactionPct" in v or "ProcessingPct" in v or "OffsiteAdsPct" in v
                ):
                    fee_formula_cells.append((ws.title, c.coordinate, v))
    assert fee_formula_cells, "Expected fee formulas referencing TransactionPct/ProcessingPct/OffsiteAdsPct"
    for sheet, coord, formula in fee_formula_cells:
        assert not _re.search(r"[^A-Za-z](0\.0[0-9]|0\.1[0-9]|0\.2[0-9]|0\.3[0-9]|0\.06[0-9])",
                               formula), f"Possible hardcoded fee value in {sheet}!{coord}: {formula}"
    print(f"PASS: {len(fee_formula_cells)} fee formula(s) reference defined names, no hardcoded fee values")

    # 5. Calculator: required formula cells present.
    for name in ("Revenue", "TotalFees", "NetProfit", "MarginPct", "RevenueNeeded", "SalePriceNeeded"):
        cell = resolve(name)
        assert isinstance(cell.value, str) and cell.value.startswith("="), (
            f"{name} missing formula, got {cell.value!r}"
        )
    print("PASS: all Profit Calculator breakdown + reverse formulas present")

    # 6. Offsite ad dropdown validations: one on Profit Calculator (single
    #    cell), one on Sales Log (100-row range).
    calc_dvs = calc.data_validations.dataValidation
    assert len(calc_dvs) == 1, f"expected 1 data validation on Profit Calculator, found {len(calc_dvs)}"
    assert calc_dvs[0].type == "list"
    log_dvs = log.data_validations.dataValidation
    assert len(log_dvs) == 1, f"expected 1 data validation on Sales Log, found {len(log_dvs)}"
    dv = log_dvs[0]
    assert dv.type == "list"
    sqref = str(dv.sqref)
    assert f"G{SL_FIRST_DATA_ROW}:G{SL_LAST_DATA_ROW}" in sqref, f"Offsite ad? dropdown range wrong: {sqref}"
    print("PASS: Offsite ad? Yes/No dropdowns present on both Profit Calculator and Sales Log")

    # 7. Sales Log: 100 rows, 5 sample rows, formulas on every row, totals row,
    #    this-month tile.
    total_rows = SL_LAST_DATA_ROW - SL_FIRST_DATA_ROW + 1
    assert total_rows == 100, total_rows
    filled = sum(1 for r in range(SL_FIRST_DATA_ROW, SL_LAST_DATA_ROW + 1)
                 if log.cell(row=r, column=2).value is not None)
    assert filled == len(SAMPLE_LOG_ROWS), f"expected {len(SAMPLE_LOG_ROWS)} filled sample rows, found {filled}"
    for r in range(SL_FIRST_DATA_ROW, SL_LAST_DATA_ROW + 1):
        for col in (8, 9):
            v = log.cell(row=r, column=col).value
            assert isinstance(v, str) and v.startswith("="), f"Sales Log!{log.cell(row=r,column=col).coordinate} missing formula"
    for col in (3, 4, 5, 6, 8, 9):
        v = log.cell(row=SL_TOTAL_ROW_NUM, column=col).value
        assert isinstance(v, str) and v.startswith("=SUM("), f"Totals row col {col} missing SUM formula, got {v!r}"
    tile_formula = log.cell(row=5, column=1).value
    assert isinstance(tile_formula, str) and tile_formula.startswith("=SUMIFS("), "This month's net tile missing SUMIFS formula"
    print(f"PASS: 100 Sales Log rows, {filled} sample rows filled, Fees/Net formulas on every row, "
          "totals row, this-month SUMIFS tile present")

    # 8. Whitelisted functions only.
    func_counter, disallowed_common, placeholder_hits = common.scan_workbook_formulas(wb)
    assert not disallowed_common, f"functions outside common.py's ALLOWED_FUNCTIONS: {disallowed_common}"
    disallowed_task = set(func_counter) - TASK_ALLOWED_FUNCTIONS
    assert not disallowed_task, f"functions outside the task brief's whitelist: {disallowed_task}"
    assert not placeholder_hits, f"leftover [placeholder] text found: {placeholder_hits}"
    print("PASS: functions used are all whitelisted:", sorted(func_counter))
    print("PASS: no leftover [placeholder] text")

    print("\nAll --verify checks passed.")


if __name__ == "__main__":
    if "--verify" in sys.argv:
        verify()
    else:
        build()
