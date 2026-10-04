#!/usr/bin/env python3
"""
Build the FREE "Simple Rental Property Analysis" lead-magnet workbook.

A deliberately thin, single-property deal analyzer: Inputs sheet + Analysis
sheet. No multi-property tracking, no dashboards, no charts -- that is what the
paid Rental Property Manager Toolkit (P5) is for. Defaults match the website's
rental-property-roi-calculator defaults so the two can be cross-checked.

Run:    .venv/bin/python build/free_rental_analysis.py
Verify: .venv/bin/python build/free_rental_analysis.py --verify   (LibreOffice recalc)
"""
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT_PATH = ROOT / "dist" / "free" / "GrowthChief-Simple-Rental-Property-Analysis.xlsx"
def _find_soffice():
    import subprocess as _sp
    for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice", shutil.which("soffice") or ""):
        if p and Path(p).exists():
            try:
                if _sp.run([p, "--version"], capture_output=True, timeout=60).returncode == 0:
                    return p
            except Exception:
                pass
    return "/Applications/LibreOffice.app/Contents/MacOS/soffice"


SOFFICE = _find_soffice()
PRODUCT_NAME = "Simple Rental Property Analysis (Free)"

ETSY_TOOLKIT_URL = "https://growthchief.etsy.com/listing/4561168525"
ETSY_CUSTOM_URL = "https://growthchief.etsy.com/listing/4564738495"
SITE_URL = "https://growthchief-tools.pages.dev/rental-property-roi-calculator/"

link_font = Font(name=common.FONT_NAME, size=11, bold=True, color=common.ACCENT_COLOR, underline="single")
section_font = Font(name=common.FONT_NAME, size=13, bold=True, color=common.HEADER_FILL_COLOR)
input_font = Font(name=common.FONT_NAME, size=11, bold=True, color="1D4ED8")

USD = '"$"#,##0;[Red]-"$"#,##0'
USD2 = '"$"#,##0.00;[Red]-"$"#,##0.00'
PCT = "0.00%;[Red]-0.00%"
PCT1 = "0.0%"

# (key, label, default, number format, note)
INPUTS = [
    ("price", "Purchase price ($)", 200000, USD, ""),
    ("closing", "Closing / purchase costs ($)", 4000, USD, "Legal fees, inspection, transfer taxes, lender fees."),
    ("rehab", "Rehab / initial repairs ($)", 0, USD, "Paid in cash. Counts as cash invested."),
    ("down", "Down payment (% of price)", 0.20, PCT1, "Enter 100% for an all-cash deal."),
    ("rate", "Loan interest rate (% APR)", 0.065, "0.00%", "Fixed rate, principal & interest only."),
    ("term", "Loan term (years)", 30, "0", ""),
    ("rent", "Monthly rent ($)", 1800, USD, ""),
    ("vac", "Vacancy allowance (% of rent)", 0.05, PCT1, "Share of the year the unit is empty or rent goes unpaid."),
    ("tax", "Property tax ($ per year)", 2400, USD, ""),
    ("ins", "Insurance ($ per year)", 1200, USD, ""),
    ("mgmt", "Property management (% of rent collected)", 0.08, PCT1, "Enter 0% if you self-manage."),
    ("maint", "Maintenance & repairs (% of rent collected)", 0.05, PCT1, ""),
    ("capex", "Capital expenditure reserve (% of rent collected)", 0.05, PCT1, "Saving for roof, HVAC, appliances."),
    ("other", "Other monthly costs ($)", 0, USD, "HOA dues, landlord-paid utilities, lawn care."),
]
FIRST_INPUT_ROW = 5
REF = {k: f"Inputs!$C${FIRST_INPUT_ROW + i}" for i, (k, *_r) in enumerate(INPUTS)}


def build_start_here(wb):
    steps = [
        "Open the Inputs tab and replace the blue example numbers with your own deal.",
        "Open the Analysis tab: payment, NOI, cash flow, cap rate, cash-on-cash, DSCR, the 1% rule and GRM calculate on their own.",
        "Change one input at a time to see what moves the result: rent, rate and down payment matter most.",
        "Percent inputs are entered as percentages (type 5% or 0.05). Dollar inputs are plain numbers.",
        "Works in Excel and Google Sheets as-is -- no macros, no add-ons.",
    ]
    settings = [{"label": "Property label", "value": "My example rental", "name": None,
                 "note": "Optional name for this deal; shows on the Analysis tab."}]
    result = common.build_start_here(
        wb, PRODUCT_NAME, steps, settings,
        guide_link_placeholder="This free version is simple by design; the download page has a full walkthrough.",
        toolkit_name="GrowthChief spreadsheet toolkits",
    )
    ws = result["ws"]
    label_ref = f"'Start Here'!$C${result['settings_range'][0]}"
    row = result["settings_range"][1] + 2
    ws.cell(row=row, column=2, value="Want more?").font = section_font
    row += 1
    c = ws.cell(row=row, column=2, value=(
        "This free analyzer screens ONE property at a time. The Rental Property Manager Toolkit "
        "tracks up to 8 properties with rent, expenses, tenants and maintenance, and a dashboard of real cash flow."))
    c.font = common.body_font
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 48
    row += 2
    l1 = ws.cell(row=row, column=2, value="See the Rental Property Manager Toolkit on Etsy")
    l1.font = link_font
    l1.hyperlink = ETSY_TOOLKIT_URL
    row += 1
    l2 = ws.cell(row=row, column=2, value="Need something custom? Custom spreadsheets, made to order ($99)")
    l2.font = link_font
    l2.hyperlink = ETSY_CUSTOM_URL
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    row += 1
    l3 = ws.cell(row=row, column=2, value="Free online rental property ROI calculator")
    l3.font = link_font
    l3.hyperlink = SITE_URL
    row += 2
    n = ws.cell(row=row, column=2, value="Estimates only, not financial advice. Check real numbers with a qualified professional.")
    n.font = common.note_font
    return label_ref


def build_inputs(wb):
    ws = common.new_data_sheet(wb, "Inputs", tab_color=common.TAB_SLATE, gridlines=False)
    common.set_col_widths(ws, {"A": 3, "B": 46, "C": 18, "D": 56})
    common.style_title(ws, "B1", "Inputs")
    ws.cell(row=2, column=2, value="Replace the blue numbers with your deal. Everything on the Analysis tab updates.").font = common.subtitle_font
    common.style_header_row(ws, 4, 2, 4, height=22)
    for col, v in enumerate(["Input", "Value", "Notes"], start=2):
        ws.cell(row=4, column=col, value=v)
    for i, (key, label, default, fmt, note) in enumerate(INPUTS):
        r = FIRST_INPUT_ROW + i
        ws.cell(row=r, column=2, value=label).font = common.label_font
        c = ws.cell(row=r, column=3, value=default)
        c.font = input_font
        c.fill = common.accent_light_fill
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
        n = ws.cell(row=r, column=4, value=note)
        n.font = common.note_font
        for col in (2, 3, 4):
            ws.cell(row=r, column=col).border = common.thin_border
    return ws


def build_analysis(wb, label_ref):
    ws = common.new_data_sheet(wb, "Analysis", tab_color=common.TAB_TEAL, gridlines=False)
    common.set_col_widths(ws, {"A": 3, "B": 46, "C": 20, "D": 60})
    common.style_title(ws, "B1", "Analysis")
    ws["B2"] = f'={label_ref}'
    ws["B2"].font = common.subtitle_font
    R = REF
    rows = [
        ("section", "Financing"),
        ("loan_amt", "Loan amount", f"={R['price']}*(1-{R['down']})", USD, "Price less down payment."),
        ("down_amt", "Down payment", f"={R['price']}*{R['down']}", USD, ""),
        ("pmt", "Monthly mortgage payment (P&I)",
         f"=IF(C{{loan_amt}}<=0,0,-PMT({R['rate']}/12,{R['term']}*12,C{{loan_amt}}))", USD2, "Calculated with PMT."),
        ("debt", "Annual debt service", "=C{pmt}*12", USD, ""),
        ("section", "Income and operating costs (per year)"),
        ("gross", "Gross potential rent", f"={R['rent']}*12", USD, ""),
        ("vacloss", "Vacancy loss", f"=C{{gross}}*{R['vac']}", USD, ""),
        ("coll", "Rent collected", "=C{gross}-C{vacloss}", USD, ""),
        ("opex", "Operating costs (before mortgage)",
         f"={R['tax']}+{R['ins']}+C{{coll}}*({R['mgmt']}+{R['maint']}+{R['capex']})+{R['other']}*12", USD,
         "Tax + insurance + management/maintenance/capex on rent collected + other."),
        ("noi", "Net operating income (NOI)", "=C{coll}-C{opex}", USD, ""),
        ("section", "Cash flow"),
        ("cfa", "Annual cash flow (after mortgage)", "=C{noi}-C{debt}", USD, ""),
        ("cfm", "Monthly cash flow", "=C{cfa}/12", USD, ""),
        ("inv", "Total cash invested", f"=C{{down_amt}}+{R['closing']}+{R['rehab']}", USD, "Down payment + closing + rehab."),
        ("section", "Return metrics"),
        ("cap", "Cap rate", f"=IF({R['price']}>0,C{{noi}}/{R['price']},0)", PCT, "NOI / purchase price. Ignores financing."),
        ("coc", "Cash-on-cash return", '=IF(C{inv}>0,C{cfa}/C{inv},"n/a")', PCT, "Annual cash flow / cash invested."),
        ("dscr", "DSCR (NOI / debt service)", '=IF(C{debt}>0,C{noi}/C{debt},"n/a (no loan)")', '0.00"x"',
         "Lenders often look for roughly 1.0 to 1.25 or more; varies by lender."),
        ("one", "Rent / (price + rehab) -- 1% rule", f"=IF({R['price']}+{R['rehab']}>0,{R['rent']}/({R['price']}+{R['rehab']}),0)", PCT,
         "Quick screen only."),
        ("onechk", "1% rule check", '=IF(C{one}>=0.01,"Meets the 1% rule","Below the 1% rule")', "@", ""),
        ("grm", "Gross rent multiplier (price / annual rent)", f'=IF(C{{gross}}>0,{R["price"]}/C{{gross}},"n/a")', '0.00"x"',
         "Lower can mean more rent per dollar of price."),
    ]
    # assign rows
    r = 4
    rowmap = {}
    for item in rows:
        if item[0] != "section":
            rowmap[item[0]] = r
        r += 1
    r = 4
    for item in rows:
        if item[0] == "section":
            ws.cell(row=r, column=2, value=item[1]).font = section_font
            r += 1
            continue
        key, label, formula, fmt, note = item
        ws.cell(row=r, column=2, value=label).font = common.label_font
        c = ws.cell(row=r, column=3, value=formula.format(**rowmap))
        c.number_format = fmt
        c.font = common.accent_font
        c.fill = common.accent_light_fill
        c.alignment = Alignment(horizontal="right")
        ws.cell(row=r, column=4, value=note).font = common.note_font
        for col in (2, 3, 4):
            ws.cell(row=r, column=col).border = common.thin_border
        r += 1
    ws.cell(row=r + 1, column=2, value=(
        "Estimates only, not financial advice. No income tax, depreciation, selling costs or rent growth are modeled."
    )).font = common.note_font
    return ws, rowmap


def build():
    wb = Workbook()
    label_ref = build_start_here(wb)
    build_inputs(wb)
    ws, rowmap = build_analysis(wb, label_ref)
    assert wb.sheetnames == ["Start Here", "Inputs", "Analysis"], wb.sheetnames
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} ({OUT_PATH.stat().st_size:,} bytes)")
    print("Analysis rows:", rowmap)
    return rowmap


def verify():
    rowmap = build.__wrapped_rowmap if hasattr(build, "__wrapped_rowmap") else None
    wb = load_workbook(OUT_PATH)
    an = wb["Analysis"]
    # no unprefixed cross-sheet refs to Inputs / Start Here: every Analysis formula's refs to
    # input cells must carry a sheet prefix, and local refs must be to Analysis cells that exist.
    for row in an.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                body = re.sub(r'"[^"]*"', "", v)
                for m in re.finditer(r"(?<![!\w$'])\$?[A-Z]{1,3}\$?\d+", body):
                    ref = m.group(0).replace("$", "")
                    tgt = an[ref].value
                    assert tgt not in (None, ""), f"{c.coordinate}: unprefixed ref {ref} points at empty Analysis cell"
    if not Path(SOFFICE).exists():
        print("LibreOffice not available; evaluating formulas with the built-in mini evaluator instead.")
        vals = mini_eval(wb)
        for r, v in sorted(vals.items()):
            print(f"  {an.cell(row=r, column=2).value:<48} {v}")
        return vals
    with tempfile.TemporaryDirectory() as td:
        subprocess.run([SOFFICE, "--headless", "--convert-to", "xlsx", "--outdir", td, str(OUT_PATH)],
                       capture_output=True, timeout=180, check=True)
        cw = load_workbook(Path(td) / OUT_PATH.name, data_only=True)["Analysis"]
        for row in cw.iter_rows(min_row=4):
            if row[1].value and row[2].value is not None:
                print(f"  {row[1].value:<48} {row[2].value}")


def mini_eval(wb):
    """Tiny independent evaluator for the Analysis formulas (IF/PMT/arithmetic only)."""
    an, inp = wb["Analysis"], wb["Inputs"]
    cache = {}

    def PMT(rate, nper, pv):
        return -pv * rate / (1 - (1 + rate) ** -nper) if rate else -pv / nper

    def ev(row):
        if row in cache:
            return cache[row]
        f = an.cell(row=row, column=3).value
        if not (isinstance(f, str) and f.startswith("=")):
            return f
        e = f[1:]
        e = re.sub(r"Inputs!\$C\$(\d+)", lambda m: repr(inp.cell(row=int(m.group(1)), column=3).value), e)
        e = re.sub(r"(?<![A-Za-z!$])C(\d+)", lambda m: repr(ev(int(m.group(1)))), e)
        e = e.replace("<>", "!=")
        e = re.sub(r"IF\(", "IFF(", e)
        v = eval(e, {"IFF": lambda c, a, b: a if c else b, "PMT": PMT})
        cache[row] = v
        return v

    for r in range(4, an.max_row + 1):
        if isinstance(an.cell(row=r, column=3).value, str) and an.cell(row=r, column=3).value.startswith("="):
            ev(r)
    return cache


if __name__ == "__main__":
    if "--verify" in sys.argv:
        verify()
    else:
        build()
